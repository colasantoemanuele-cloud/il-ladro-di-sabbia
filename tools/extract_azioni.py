#!/usr/bin/env python3
"""
Estrae la tabella delle azioni dal foglio "Azioni" di
ladro_di_sabbia_bilanciamento.xlsx e la scrive come data/azioni.json,
il formato letto a runtime da ActionDatabase (scripts/data/action_database.gd).

Uso: python3 tools/extract_azioni.py
Richiede: pip install openpyxl

Rilanciare questo script dopo ogni modifica al foglio "Azioni" per
tenere data/azioni.json sincronizzato con l'Excel (fonte di verita').
"""

import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX_PATH = ROOT / "ladro_di_sabbia_bilanciamento.xlsx"
JSON_PATH = ROOT / "data" / "azioni.json"

HEADER_ROW = 5
FIRST_DATA_ROW = 6


## Estrae le azioni dal workbook già aperto, senza scrivere nulla su disco:
## riusata sia da main() (scrittura reale) sia da tools/check_data_consistency.py
## (confronto Excel vs JSON già scritto, senza duplicare la logica di
## parsing — un'unica fonte di verità su COME si legge il foglio Azioni).
## Restituisce (lista azioni, salario_mediano, ore_anno).
def estrai_azioni(wb) -> tuple:
    ws = wb["Azioni"]

    salario_mediano = ws["B2"].value
    ore_anno = ws["B3"].value

    actions = []
    for row in ws.iter_rows(min_row=FIRST_DATA_ROW, max_row=ws.max_row):
        nome = row[0].value
        categoria = row[1].value
        if nome is None or categoria is None:
            # Riga vuota o nota istruttiva in fondo al foglio: non e' un'azione.
            continue
        valore_eur = row[3].value
        # Le colonne E/F sono formule Excel (=D/$B$2, =E/$B$3): calcolate qui
        # direttamente da D invece di leggere il valore cache di openpyxl
        # (data_only=True legge solo l'ULTIMO valore calcolato da Excel/
        # LibreOffice al salvataggio — righe aggiunte via script, come quelle
        # della Fase 10, non hanno mai una cache e leggerebbero None). Più
        # robusto: funziona sempre, a prescindere da come il file e' stato
        # salvato l'ultima volta.
        effetto_ore = valore_eur / salario_mediano
        effetto_anni = effetto_ore / ore_anno
        actions.append({
            "nome": nome,
            "categoria": categoria,
            "costo_tempo_figlia_ore": row[2].value,
            "valore_economico_eur": valore_eur,
            "effetto_sabbia_padre_ore": effetto_ore,
            "effetto_sabbia_padre_anni": effetto_anni,
            "rischio_pct": row[6].value,
            "moralita": row[7].value,
            "unica_per_run": row[8].value == "VERO",
            "fonte_nota": row[9].value,
        })

    return actions, salario_mediano, ore_anno


def main() -> None:
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    actions, salario_mediano, ore_anno = estrai_azioni(wb)

    output = {
        "_meta": {
            "fonte": "ladro_di_sabbia_bilanciamento.xlsx, foglio Azioni",
            "salario_mediano_eur_ora": salario_mediano,
            "ore_per_anno": ore_anno,
            "numero_azioni": len(actions),
        },
        "azioni": actions,
    }

    JSON_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(JSON_PATH, "w", encoding="utf-8") as f:
        json.dump(output, f, ensure_ascii=False, indent=2)

    print(f"Scritte {len(actions)} azioni in {JSON_PATH}")
    if len(actions) != 62:
        print(
            f"ATTENZIONE: attese 62 azioni (60 core confermate + 2 aggiunte in Fase 10 per far "
            f"scendere Rivalita' Criminale/Fama Pubblica), ma il foglio Excel ne contiene {len(actions)}.",
            file=sys.stderr,
        )
    n_uniche = sum(1 for a in actions if a["unica_per_run"])
    print(f"Di cui marcate 'Unica per run': {n_uniche}")


if __name__ == "__main__":
    main()
