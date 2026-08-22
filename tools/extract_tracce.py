#!/usr/bin/env python3
"""
Fase 9a — estrae le 10 tracce (7 normali + 3 bonus) dal foglio "Endgame" di
ladro_di_sabbia_bilanciamento.xlsx e le scrive come data/tracce.json, letto
a runtime da TrackDatabase (scripts/data/track_database.gd), sullo stesso
pattern di tools/extract_azioni.py per le 60 azioni core.

Uso: python3 tools/extract_tracce.py
Richiede: pip install openpyxl

Rilanciare dopo ogni modifica al foglio "Endgame" per tenere
data/tracce.json sincronizzato con l'Excel (fonte di verita').
"""

import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX_PATH = ROOT / "ladro_di_sabbia_bilanciamento.xlsx"
JSON_PATH = ROOT / "data" / "tracce.json"

TRACCE_NORMALI_FIRST_ROW = 15
TRACCE_NORMALI_LAST_ROW = 28
TRACCE_BONUS_FIRST_ROW = 34
TRACCE_BONUS_LAST_ROW = 36

RANGO_RE = re.compile(r"Rango\s+(\d+)\s*-\s*(.+)")


def main() -> None:
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb["Endgame"]

    tracce_normali = []
    for row in ws.iter_rows(min_row=TRACCE_NORMALI_FIRST_ROW, max_row=TRACCE_NORMALI_LAST_ROW):
        traccia = row[0].value
        rango_raw = row[1].value
        if traccia is None or rango_raw is None:
            continue
        m = RANGO_RE.match(rango_raw)
        if not m:
            print(f"ATTENZIONE: formato rango non riconosciuto: {rango_raw!r}", file=sys.stderr)
            continue
        tracce_normali.append({
            "traccia": traccia,
            "rango": int(m.group(1)),
            "nome_rango": m.group(2).strip(),
            "costo_tempo_figlia_ore": row[2].value,
            "effetto_sabbia_padre_ore": row[3].value,
            "rischio_pct": row[4].value,
            "nota": row[5].value,
        })

    tracce_bonus = []
    for row in ws.iter_rows(min_row=TRACCE_BONUS_FIRST_ROW, max_row=TRACCE_BONUS_LAST_ROW):
        nome = row[0].value
        descrizione = row[1].value
        if nome is None:
            continue
        tracce_bonus.append({
            "traccia": nome,
            "descrizione": descrizione,
        })

    output = {
        "_meta": {
            "fonte": "ladro_di_sabbia_bilanciamento.xlsx, foglio Endgame, sezioni 2 e 3",
            "numero_tracce_normali_righe": len(tracce_normali),
            "numero_tracce_bonus": len(tracce_bonus),
        },
        "tracce_normali": tracce_normali,
        "tracce_bonus": tracce_bonus,
    }

    JSON_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(JSON_PATH, "w", encoding="utf-8") as f:
        json.dump(output, f, ensure_ascii=False, indent=2)

    n_tracce_uniche = len(set(t["traccia"] for t in tracce_normali))
    print(f"Scritte {len(tracce_normali)} righe rango ({n_tracce_uniche} tracce normali) e {len(tracce_bonus)} tracce bonus in {JSON_PATH}")
    if len(tracce_normali) != 14:
        print(f"ATTENZIONE: attese 14 righe (7 tracce x 2 ranghi), trovate {len(tracce_normali)}.", file=sys.stderr)
    if n_tracce_uniche != 7:
        print(f"ATTENZIONE: attese 7 tracce normali distinte, trovate {n_tracce_uniche}.", file=sys.stderr)
    if len(tracce_bonus) != 3:
        print(f"ATTENZIONE: attese 3 tracce bonus, trovate {len(tracce_bonus)}.", file=sys.stderr)


if __name__ == "__main__":
    main()
